"""文本级版本差异。不判断两版方案在业务含义上是否相同。"""

from __future__ import annotations

import difflib

from docx import Document
from openpyxl import load_workbook

from app.config import Settings
from app.db.models import ArtifactVersion
from app.services.document_service import safe_storage_path


def artifact_plain_text(settings: Settings, artifact: ArtifactVersion) -> str:
    path = safe_storage_path(settings.artifact_dir, artifact.relative_path)
    if artifact.artifact_type == "proposal_docx":
        document = Document(path)
        lines = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(lines)
    if artifact.artifact_type == "plan_xlsx":
        workbook = load_workbook(path, data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                lines.append("\t".join("" if value is None else str(value) for value in row))
        return "\n".join(lines)
    return path.read_text(encoding="utf-8")


def unified_text_diff(left_label: str, left_text: str, right_label: str, right_text: str) -> str:
    return "".join(
        difflib.unified_diff(
            left_text.splitlines(keepends=True),
            right_text.splitlines(keepends=True),
            fromfile=left_label,
            tofile=right_label,
        )
    )


