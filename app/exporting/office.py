"""把项目书和任务计划写成可再次打开的 Office 文件。待确认内容保留并醒目标出。"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DISCLAIMER = "本文件由演示用户在本地确认，不是正式发布的业务文件，也不能代替业务审批。"

_PLAN_COLUMNS = (
    ("task_id", "任务编号", 14),
    ("title", "任务名称", 28),
    ("description", "任务描述", 36),
    ("owner_role", "责任角色", 16),
    ("start_date", "开始日期", 14),
    ("end_date", "结束日期", 14),
    ("dependency_ids", "依赖任务", 16),
    ("deliverables", "交付物", 24),
    ("acceptance_criteria", "验收标准", 24),
    ("risk_notes", "风险", 24),
    ("status", "状态", 16),
)


def build_proposal_docx(
    proposal: dict,
    tasks: list[dict],
    evidence_items: list[dict],
    *,
    prompt_version: str | None,
) -> bytes:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    document.add_heading(str(proposal.get("title") or "项目书草稿"), 0)
    _warning(document, DISCLAIMER)
    if proposal.get("mock"):
        _warning(document, "生成时使用了 mock 模型，不能当作真实模型输出。")
    document.add_paragraph(f"提示词版本：{prompt_version or '未记录'}")
    document.add_paragraph("审核状态：demo_approved。这只表示演示用户确认了草稿。")

    sections = list(proposal.get("sections") or [])
    citations = _citation_index(sections, tasks, evidence_items)
    document.add_heading("目录", level=1)
    if not sections:
        document.add_paragraph("没有章节。空内容已保留，没有用默认正文填补。")
    for index, section in enumerate(sections, start=1):
        document.add_paragraph(f"{index}. {section.get('heading') or '未命名章节'}")

    document.add_heading("正文", level=1)
    for section in sections:
        document.add_heading(str(section.get("heading") or "未命名章节"), level=2)
        body = str(section.get("body") or "")
        if section.get("needs_confirmation"):
            body = f"【待确认】{body}"
        marks = _marks(section.get("evidence_ids") or [], citations)
        paragraph = document.add_paragraph(f"{body} {marks}".strip())
        if section.get("needs_confirmation"):
            _mark_unconfirmed(paragraph)

    document.add_heading("任务摘要", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, name in zip(
        table.rows[0].cells, ("任务编号", "任务名称", "责任角色", "状态"), strict=True
    ):
        cell.text = name
    for task in tasks:
        row = table.add_row().cells
        title = str(task.get("title") or "")
        pending = bool(task.get("needs_confirmation"))
        if pending:
            title = f"【待确认】{title}"
        row[0].text = str(task.get("task_id") or "")
        row[1].text = title
        row[2].text = str(task.get("owner_role") or "")
        row[3].text = "待确认" if pending else "演示确认草稿"
        description = str(task.get("description") or "")
        if description:
            document.add_paragraph(f"{task.get('task_id') or ''} {description}")

    document.add_heading("引用来源附录", level=1)
    document.add_paragraph("引用编号只说明草稿使用了该片段，不代表内容已经过业务批准。")
    if not citations:
        document.add_paragraph("本次没有可定位的证据编号。待确认内容已保留，没有被删除。")
    else:
        appendix = document.add_table(rows=1, cols=5)
        appendix.style = "Table Grid"
        for cell, name in zip(
            appendix.rows[0].cells,
            ("编号", "来源", "页码", "章节", "片段"),
            strict=True,
        ):
            cell.text = name
        for entry in citations.values():
            item = entry["item"] or {}
            row = appendix.add_row().cells
            row[0].text = f"[{entry['number']}]"
            row[1].text = str(item.get("title") or "未命名来源")
            row[2].text = "" if item.get("page_number") is None else str(item.get("page_number"))
            row[3].text = str(item.get("section") or "")
            row[4].text = str(item.get("quote") or "")[:300]

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_plan_xlsx(
    tasks: list[dict],
    *,
    title: str,
    prompt_version: str | None,
) -> bytes:
    workbook = Workbook()
    note = workbook.active
    note.title = "说明"
    note["A1"] = DISCLAIMER
    note["A2"] = "本表不是正式批准计划。状态只使用“待确认”或“演示确认草稿”。"
    note["A3"] = f"项目书：{title or '未命名'}；提示词版本：{prompt_version or '未记录'}"
    note["A1"].font = Font(bold=True, color="9C0006")
    note.column_dimensions["A"].width = 88

    sheet = workbook.create_sheet("任务计划", 0)
    headers = [label for _, label, _ in _PLAN_COLUMNS]
    for column, label in enumerate(headers, start=1):
        cell = sheet.cell(1, column, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    for row_index, task in enumerate(tasks, start=2):
        pending = bool(task.get("needs_confirmation"))
        for column, (key, _label, _width) in enumerate(_PLAN_COLUMNS, start=1):
            cell = sheet.cell(row_index, column)
            if key == "status":
                cell.value = "待确认" if pending else "演示确认草稿"
            elif key in {"start_date", "end_date"}:
                parsed = _as_date(task.get(key))
                if parsed is not None:
                    cell.value = parsed
                    cell.number_format = "YYYY-MM-DD"
                elif task.get(key):
                    cell.value = str(task.get(key))
            elif key == "dependency_ids":
                cell.value = "、".join(str(item) for item in task.get(key) or [])
            elif key == "title" and pending:
                cell.value = f"【待确认】{task.get('title') or ''}"
            else:
                cell.value = "" if task.get(key) is None else str(task.get(key))
    for column, (_key, _label, width) in enumerate(_PLAN_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(_PLAN_COLUMNS))}{max(len(tasks) + 1, 1)}"

    milestones = workbook.create_sheet("里程碑")
    milestones.append(["任务编号", "名称", "结束日期", "交付物"])
    risks = workbook.create_sheet("风险台账")
    risks.append(["任务编号", "名称", "风险"])
    for task in tasks:
        end = _as_date(task.get("end_date"))
        milestones.append(
            [
                str(task.get("task_id") or ""),
                str(task.get("title") or ""),
                end
                if end is not None
                else (str(task.get("end_date")) if task.get("end_date") else None),
                str(task.get("deliverables") or ""),
            ]
        )
        if str(task.get("risk_notes") or "").strip():
            risks.append(
                [
                    str(task.get("task_id") or ""),
                    str(task.get("title") or ""),
                    str(task.get("risk_notes")),
                ]
            )
    for extra in (milestones, risks):
        extra.freeze_panes = "A2"
        extra.auto_filter.ref = extra.dimensions
        if (
            extra.max_row >= 2
            and extra["C2"].value is not None
            and isinstance(extra["C2"].value, date)
        ):
            extra["C2"].number_format = "YYYY-MM-DD"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _warning(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor(0x9C, 0x00, 0x06)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _mark_unconfirmed(paragraph) -> None:
    for run in paragraph.runs:
        run.bold = True
        run.font.color.rgb = RGBColor(0x9A, 0x34, 0x00)


def _citation_index(
    sections: list[dict], tasks: list[dict], evidence_items: list[dict]
) -> dict[str, dict]:
    by_id: dict[str, dict] = {}
    for item in evidence_items:
        for key in (item.get("evidence_id"), item.get("chunk_id")):
            if key:
                by_id[str(key)] = item
    ordered: list[str] = []
    seen: set[str] = set()
    groups = list(sections)
    groups.extend(tasks)
    for group in groups:
        for evidence_id in group.get("evidence_ids") or []:
            token = str(evidence_id)
            if token and token not in seen:
                seen.add(token)
                ordered.append(token)
    return {
        token: {"number": index, "item": by_id.get(token)}
        for index, token in enumerate(ordered, start=1)
    }


def _marks(evidence_ids: list, citations: dict[str, dict]) -> str:
    numbers = []
    for evidence_id in evidence_ids:
        entry = citations.get(str(evidence_id))
        if entry is not None:
            numbers.append(f"[{entry['number']}]")
    return " ".join(numbers)


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and len(value) >= 10:
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return None
    return None
