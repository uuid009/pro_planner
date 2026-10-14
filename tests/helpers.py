from io import BytesIO

import fitz
from docx import Document
from openpyxl import Workbook


def make_pdf(text: str) -> bytes:
    document = fitz.open()
    page = document.new_page()
    # china-s 是 PyMuPDF 内置中文字体；默认 Helvetica 会把汉字写成占位符。
    page.insert_text((72, 72), text, fontname="china-s", fontsize=12)
    payload = document.tobytes()
    document.close()
    return payload


def make_empty_pdf() -> bytes:
    document = fitz.open()
    document.new_page()
    payload = document.tobytes()
    document.close()
    return payload


def make_docx(
    *paragraphs: str, heading: str | None = None, table_rows: list[list[str]] | None = None
) -> bytes:
    document = Document()
    if heading:
        document.add_heading(heading, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx(rows: list[list[str]], sheet: str = "计划") -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for row in rows:
        worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
