from datetime import date
from io import BytesIO

from docx import Document
from openpyxl import load_workbook

from app.exporting.office import build_plan_xlsx, build_proposal_docx


def test_docx_keeps_unconfirmed_text_and_reopens():
    raw = build_proposal_docx(
        {
            "title": "示范项目书",
            "mock": True,
            "sections": [
                {
                    "heading": "目标",
                    "body": "完成一期上线",
                    "evidence_ids": ["e1"],
                    "needs_confirmation": True,
                }
            ],
        },
        [
            {
                "task_id": "T1",
                "title": "确认资料",
                "description": "保留待确认描述",
                "owner_role": "策划专责",
                "needs_confirmation": True,
            }
        ],
        [
            {
                "evidence_id": "e1",
                "title": "台账说明",
                "page_number": 3,
                "section": "背景",
                "quote": "原文片段",
            }
        ],
        prompt_version="phase3-v1",
    )
    document = Document(BytesIO(raw))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            text += "\n" + "\t".join(cell.text for cell in row.cells)
    assert "【待确认】完成一期上线" in text
    assert "【待确认】确认资料" in text
    assert "不是正式发布" in text
    assert "引用来源附录" in text
    assert "[1]" in text
    assert "mock 模型" in text
    assert "原文片段" in text


def test_xlsx_uses_excel_dates_and_does_not_claim_formal_approval():
    raw = build_plan_xlsx(
        [
            {
                "task_id": "T1",
                "title": "确认资料",
                "description": "整理目标",
                "owner_role": "策划专责",
                "start_date": "2026-01-01",
                "end_date": "2026-01-15",
                "dependency_ids": [],
                "deliverables": "确认单",
                "acceptance_criteria": "书面确认",
                "risk_notes": "资料不足",
                "needs_confirmation": True,
            },
            {
                "task_id": "T2",
                "title": "编制方案",
                "description": "写项目书",
                "owner_role": "实施专责",
                "start_date": "下周一",
                "end_date": "2026-06-30",
                "dependency_ids": ["T1"],
                "deliverables": "项目书",
                "acceptance_criteria": "人工审核",
                "risk_notes": "",
                "needs_confirmation": False,
            },
        ],
        title="示范项目书",
        prompt_version="phase3-v1",
    )
    workbook = load_workbook(BytesIO(raw))
    sheet = workbook["任务计划"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref.startswith("A1:")
    assert sheet["E2"].value.date() == date(2026, 1, 1)
    assert sheet["E2"].number_format == "YYYY-MM-DD"
    assert sheet["E3"].value == "下周一"
    assert sheet["A2"].value == "T1"
    assert "【待确认】" in sheet["B2"].value
    text = " ".join(
        str(value)
        for row in workbook["说明"].iter_rows(values_only=True)
        for value in row
        if value is not None
    )
    text += " " + " ".join(str(cell.value) for cell in sheet[2] if cell.value is not None)
    assert "不是正式批准计划" in text
    assert "已正式批准" not in text
    assert workbook["里程碑"].freeze_panes == "A2"
    assert workbook["风险台账"]["A2"].value == "T1"
